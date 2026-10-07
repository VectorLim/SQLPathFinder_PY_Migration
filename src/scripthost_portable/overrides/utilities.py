"""Small operation overrides; original ScriptHost owns task semantics."""

import SPFLib.SPFUtilities.utils as legacy


class Utilities(legacy.Utilities):
    def _run_robocopy(self, command, arguments, pass_codes, error_codes):
        if legacy.os.name == 'nt':
            return super()._run_robocopy(command, arguments, pass_codes, error_codes)
        from scripthost_portable.file_operations import robocopy_files

        # Original caller supplies quoted paths/patterns, followed by /R and /W.
        retry_index = next((index for index, value in enumerate(arguments) if value.startswith('/R:')))
        source, destination = (value[1:-1] for value in arguments[:2])
        patterns = [value[1:-1] for value in arguments[2:retry_index]]
        try:
            code = robocopy_files(
                source, destination, patterns,
                int(arguments[retry_index][3:]), int(arguments[retry_index + 1][3:]),
                arguments[retry_index + 2:],
            )
        except (OSError, ValueError) as error:
            self.Console(str(error))
            raise legacy.SPFCMDRunExitWithErrorCodeException(str(error), 16, False) from error
        return (code in pass_codes, code)

    def _move_file(self, srcFile, DstFile):
        if legacy.os.name == 'nt':
            return super()._move_file(srcFile, DstFile)
        try:
            legacy.shutil.move(srcFile, DstFile)
        except OSError as error:
            raise legacy.SPFCMDRunExitWithErrorCodeException(str(error), 1, False) from error
        return (True, 0)

    def setEnv(self, EnvVarName, EvnVarValue):
        if legacy.os.name != 'nt':
            for name in [key for key in legacy.os.environ if key.upper() == EnvVarName.upper()]:
                del legacy.os.environ[name]
            legacy.os.environ[EnvVarName.upper()] = str(EvnVarValue)
        return super().setEnv(EnvVarName, EvnVarValue)

    def unzipString(self, inputStringToDeCompress):
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        try:
            self.logger.info('{0} - len(inputStringToDeCompress) : {1}'.format(calling_func, len(inputStringToDeCompress)))
            if legacy.isPYTHON2:
                decompressedOutputString = legacy.zlib.decompress(legacy.base64.standard_b64decode(inputStringToDeCompress), legacy.zlib.MAX_WBITS | 32).replace('\r\n', '\n')
            else:
                __t = legacy.zlib.decompress(legacy.base64.standard_b64decode(inputStringToDeCompress), legacy.zlib.MAX_WBITS | 32)
                try:
                    decompressedOutputString = __t.decode('utf-8').replace('\r\n', '\n')
                except UnicodeDecodeError:
                    decompressedOutputString = __t.decode(encoding=self.detectCharacterEncoding(__t)).replace('\r\n', '\n')
            del inputStringToDeCompress
            self.logger.info('{0} - len(decompressedOutputString) : {1}'.format(calling_func, len(decompressedOutputString)))
            return decompressedOutputString
        except Exception as err:
            self.logger.warn('{0} - Error in unzipString : {1}'.format(calling_func, err))
            return inputStringToDeCompress

    def Run_R(self, MyMode, Command1, MyLocal, WorkDir, ll_AppSvr, MyShowSQL=False, MyShowOut=False, MyArgs='', MyOS='W'):
        if legacy.os.name != 'nt' and WorkDir == '.\\':
            WorkDir = '.'
        return super().Run_R(MyMode, Command1, MyLocal, WorkDir, ll_AppSvr, MyShowSQL, MyShowOut, MyArgs, MyOS)

    def _web_auth_type(self):
        if legacy.os.name == 'nt':
            return super()._web_auth_type()
        return None

    def _web_verify(self):
        if legacy.os.name == 'nt':
            return super()._web_verify()
        return True

    def _file_pattern_separator(self):
        if legacy.os.name == 'nt':
            return super()._file_pattern_separator()
        return legacy.os.sep

    def SPFDelete(self, MySrc, quietMode=True, forceDelete=True, displayPrompt=True):
        if legacy.os.name == 'nt':
            return super().SPFDelete(MySrc, quietMode, forceDelete, displayPrompt)
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        self.logger.debug('{0} - To delete MySrc: {1}'.format(calling_func, MySrc))
        self.logger.debug('{0} - quietMode: {1}'.format(calling_func, quietMode))
        self.logger.debug('{0} - forceDelete: {1}'.format(calling_func, forceDelete))
        try:
            if displayPrompt is True:
                self.ConsoleWithTimeStamp('Starting Delete Applet, v2.5')
            if self.IsEmptyOrNone(MySrc) is True:
                self.Console('No delete files specified ...')
                return
            rdr = legacy.csv.reader(legacy.StringIO(MySrc), delimiter=',', skipinitialspace=True)
            MySrc = ','.join(['"{0}"'.format(legacy.re.sub('<c>', ',', item1.strip('"'), legacy.re.IGNORECASE)) for item1 in next(rdr) if item1.strip('\'" ') != ''])
            self.logger.debug('{0} - parsed MySrc: {1}'.format(calling_func, MySrc))
            if self.IsEmptyOrNone(MySrc) is True:
                self.Console('No delete files specified ...')
                return
            from scripthost_portable.file_operations import delete_files
            delete_files(next(legacy.csv.reader(legacy.StringIO(MySrc), skipinitialspace=True)), forceDelete)
            if displayPrompt:
                self.ConsoleDoneWithoutTimeStamp()
            return
        except Exception as err:
            self.logger.exception('{0} - {1}'.format(calling_func, err.args[0]))
            raise

    def ConvertDLM(self, InFile, OutFile, CvtExe, IsQuiet=True):
        if legacy.os.name == 'nt':
            return super().ConvertDLM(InFile, OutFile, CvtExe, IsQuiet)
        from scripthost_portable.file_operations import clean_delimited_file
        clean_delimited_file(InFile, OutFile, self.GetFileDLM(InFile))
        return

    def IntelWW(self, date):
        if legacy.os.name == 'nt':
            return super().IntelWW(date)
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        date = date.date()
        yyyy = date.year
        ww = None
        jan1 = None
        dow = None
        try:
            if date.month == 12:
                jan1 = date.day - 32
                dow = date.isoweekday()
                if dow < 7:
                    jan1 = jan1 - dow
                if jan1 > -7:
                    yyyy = yyyy + 1
                    ww = 1
                self.logger.debug("{0} - date.month == 12 : ww: '{1}'".format(calling_func, ww))
            if ww is None:
                jan1 = legacy.datetime(yyyy, 1, 1).date()
                dow = jan1.isoweekday()
                if dow < 7:
                    jan1 = jan1 - legacy.timedelta(days=dow)
                ww = (date - jan1).days / 7 + 1
                self.logger.debug("{0} - ww: '{1}'".format(calling_func, ww))
            intelww_output = int(yyyy * 100 + ww)
            self.logger.debug("{0} - intelww_output: '{1}'".format(calling_func, intelww_output))
            return str(intelww_output)
        except Exception as err:
            self.logger.exception('{0} - {1}'.format(calling_func, err))
            raise

    def SPFEmail(self, MyLocal, MyCSVFile, MailToIn, Subject, BodyF, MailCC,
                 MailBCC, MyRole, OnlyIntel, ll_Outlook, EmailUtility='SA'):
        if legacy.os.name == 'nt':
            return super().SPFEmail(
                MyLocal, MyCSVFile, MailToIn, Subject, BodyF, MailCC, MailBCC,
                MyRole, OnlyIntel, ll_Outlook, EmailUtility,
            )
        if self.IsEmptyOrNone(MyRole) is False:
            raise RuntimeError(
                'UNRESOLVED: Linux role verification (verifyrole.exe) is unavailable; '
                'original role restriction must be retained.'
            )
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        self.logger.debug("{0} - MyLocal: '{1}'".format(calling_func, MyLocal))
        self.logger.debug("{0} - MyCSVFile: '{1}'".format(calling_func, MyCSVFile))
        self.logger.debug("{0} - MailToIn: '{1}'".format(calling_func, MailToIn))
        self.logger.debug("{0} - Subject: '{1}'".format(calling_func, Subject))
        self.logger.debug("{0} - BodyF: '{1}'".format(calling_func, BodyF))
        self.logger.debug("{0} - MailCC: '{1}'".format(calling_func, MailCC))
        self.logger.debug("{0} - MailBCC: '{1}'".format(calling_func, MailBCC))
        self.logger.debug("{0} - MyRole: '{1}'".format(calling_func, MyRole))
        self.logger.debug("{0} - OnlyIntel: '{1}'".format(calling_func, OnlyIntel))
        self.logger.debug("{0} - ll_Outlook: '{1}'".format(calling_func, ll_Outlook))
        self.logger.debug("{0} - EmailUtility: '{1}'".format(calling_func, EmailUtility))
        Body = cBody = 'Please see attached files...'
        IsHTM = False
        MyCSVFileList = []
        useSMTPAuth = False
        try:
            self.Console('Starting Email Applet v4.4...')
            if self.IsEmptyOrNone(Subject) is True:
                Subject = 'Email from SQLPathFinder - 2'
                self.logger.debug("{0} - Subject: '{1}'".format(calling_func, Subject))
            if BodyF == 'Y':
                BodyF = ''
                self.logger.debug("{0} - BodyF: '{1}'".format(calling_func, BodyF))
            if self.SHisSHEntry is True:
                ll_Outlook = False
                EmailUtility = 'S'
                useSMTPAuth = False
                self.logger.debug("{0} - updated SH ll_Outlook: '{1}'".format(calling_func, ll_Outlook))
                self.logger.debug("{0} - updated SH EmailUtility: '{1}'".format(calling_func, EmailUtility))
                self.logger.debug("{0} - updated SH useSMTPAuth: '{1}'".format(calling_func, useSMTPAuth))
            if EmailUtility == 'SA':
                ll_Outlook = False
                useSMTPAuth = True
                self.logger.debug("{0} - updated SA ll_Outlook: '{1}'".format(calling_func, ll_Outlook))
                self.logger.debug("{0} - updated SA useSMTPAuth: '{1}'".format(calling_func, useSMTPAuth))
            elif EmailUtility == 'S':
                ll_Outlook = False
                useSMTPAuth = False
                self.logger.debug("{0} - updated S ll_Outlook: '{1}'".format(calling_func, ll_Outlook))
                self.logger.debug("{0} - updated S useSMTPAuth: '{1}'".format(calling_func, useSMTPAuth))
            ll_Outlook = False
            if self.IsEmptyOrNone(MailToIn) is True and self.IsEmptyOrNone(MailCC) is True and (self.IsEmptyOrNone(MailBCC) is True):
                self.Cons80()
                self.Console('The Email TO Addresses are missing ...\nExiting ...')
                self.Cons80()
            if self.IsEmptyOrNone(BodyF) is False:
                self.Console('Testing for existence of file with Email Contents ...')
                BodyF = self.Substitute_Std_Tokens(BodyF, '1')
                if legacy.os.path.exists(BodyF) is True:
                    self.Console(' File exists. Reading Contents ...\n')
                    with open(BodyF, 'r', encoding=self.detectFileEncoding(BodyF, readall=True)) as BodyContentReader:
                        Body = BodyContentReader.read()
                    if Body.upper().find('<HTML') != -1:
                        IsHTM = True
                        Body = Body.replace('<img src="gfx/image', '<img src="image')
                        Body = Body.replace('<img src=".\\gfx\\image', '<img src="image')
                else:
                    self.Console('  File does not exist. Assigning a Default Message...\n')
                    Body = cBody
            self.logger.debug("{0} - IsHTM: '{1}'".format(calling_func, IsHTM))
            if self.IsEmptyOrNone(MyCSVFile) is False:
                MyCSVFile = self.Substitute_Std_Tokens(MyCSVFile, '1')
                MyCSVFileList = [item.strip() for item in MyCSVFile.split(',')]
            elif Body == cBody:
                Body = ''
            self.logger.debug("{0} - len(MyCSVFileList): '{1}'".format(calling_func, len(MyCSVFileList)))
            self.logger.debug("{0} - MyCSVFileList: '{1}'".format(calling_func, MyCSVFileList))
            if self.IsEmptyOrNone(Subject) is False:
                Subject = self.Substitute_Std_Tokens(Subject, '1')
            userEmailAddress = ''
            from datasyncx.utils.config import get_config_value
            userEmailAddress = (get_config_value('scripthost_user_email') or '').strip()
            if not userEmailAddress:
                self.logger.warning("{0} - SCRIPTHOST_USER_EMAIL not set; 'self' recipients skipped".format(calling_func))
            MailToF = []
            ctr = 0
            if self.IsEmptyOrNone(MailToIn) is False:
                MailToF, ctr, GetEmailAdss_ret = self.GetEmailAdss(MailToIn.strip(), 'TO', userEmailAddress, MyRole, MailToF, OnlyIntel, ctr, MyLocal)
                if GetEmailAdss_ret == 'Y':
                    return
                else:
                    MailToIn = MailToF
            else:
                MailToIn = []
            MailToF = []
            if self.IsEmptyOrNone(MailCC) is False:
                MailToF, ctr, GetEmailAdss_ret = self.GetEmailAdss(MailCC, 'CC', userEmailAddress, MyRole, MailToF, OnlyIntel, ctr, MyLocal)
                if GetEmailAdss_ret == 'Y':
                    return
                else:
                    MailCC = MailToF
            else:
                MailCC = MailToF
            MailToF = []
            if self.IsEmptyOrNone(MailBCC) is False:
                MailToF, ctr, GetEmailAdss_ret = self.GetEmailAdss(MailBCC, 'BCC', userEmailAddress, MyRole, MailToF, OnlyIntel, ctr, MyLocal)
                if GetEmailAdss_ret == 'Y':
                    return
                else:
                    MailBCC = MailToF
            else:
                MailBCC = MailToF
            MailToIn, MailCC, MailBCC = ([m for m in lst if m] for lst in (MailToIn, MailCC, MailBCC))
            if not (MailToIn or MailCC or MailBCC):
                self.ConsoleWithCons80('No resolvable email addresses ...\nSkipping email ...')
                self.logger.warning('{0} - no resolvable recipients; email skipped'.format(calling_func))
                return
            if ctr == 0:
                self.Cons80()
                self.Console('There are no valid email addresses ...\nExiting ...')
                self.Cons80()
            else:
                smtpObj = None
                try:
                    ctype = 'application/octet-stream'
                    maintype, subtype = ctype.split('/', 1)
                    contentType = 'plain'
                    if IsHTM is True:
                        contentType = 'html'
                    if len(MyCSVFileList) > 0:
                        emailMessage = legacy.MIMEMultipart()
                        emailMessage.attach(legacy.MIMEText(Body, contentType, 'utf-8'))
                    else:
                        emailMessage = legacy.MIMEText(Body, contentType, 'utf-8')
                    emailMessage['Subject'] = Subject
                    emailMessage['From'] = 'atmanalytic@intel.com'
                    emailMessage['To'] = ','.join(MailToIn)
                    if len(MailCC) > 0:
                        emailMessage['Cc'] = ','.join(MailCC)
                    if len(MailBCC) > 0:
                        emailMessage['Bcc'] = ','.join(MailBCC)
                    for csvFileItem in MyCSVFileList:
                        csvFileItemPath, csvFileItemName = legacy.os.path.split(csvFileItem)
                        self.logger.debug("{0} - csvFileItem: '{1}'".format(calling_func, csvFileItem))
                        self.logger.debug("{0} - csvFileItemPath: '{1}'".format(calling_func, csvFileItemPath))
                        self.logger.debug("{0} - csvFileItemName: '{1}'".format(calling_func, csvFileItemName))
                        csvFileItem_exists, csvFileItem_modstr = self.File_Exists_Retry(csvFileItem, NameError, 3, 5, showConsoleMsgs=False)
                        if csvFileItem_exists is False:
                            errMsg = 'Error: Error while opening file : {0}'.format(csvFileItem)
                            raise Exception(errMsg)
                        if legacy.os.path.isfile(csvFileItem) is False:
                            errMsg = 'Error: Is not a file : {0}'.format(csvFileItem)
                            raise Exception(errMsg)
                        attachObj = legacy.MIMEBase(maintype, subtype)
                        with open(csvFileItem, 'rb') as attachementFileToRead:
                            attachObj.set_payload(attachementFileToRead.read())
                        legacy.encoders.encode_base64(attachObj)
                        attachObj.add_header('Content-Disposition', 'attachment', filename=csvFileItemName)
                        emailMessage.attach(attachObj)
                    self.logger.debug("{0} - useSMTPAuth : '{1}'".format(calling_func, useSMTPAuth))
                    from datasyncx.utils.send_mail import get_smtp_service_module
                    del emailMessage['Bcc']
                    smtpObj = get_smtp_service_module().get_smtp_service()
                    try:
                        smtpObj.sendmail(emailMessage['From'], MailToIn + MailCC + MailBCC, emailMessage.as_string())
                    finally:
                        smtpObj.quit()
                    self.ConsoleDoneWithTimeStamp()
                    return
                except Exception as err:
                    self.logger.exception('{0} - {1}'.format(calling_func, err.args[0]))
                    raise
                finally:
                    smtpObj = None
            self.ConsoleDoneWithTimeStamp()
        except Exception as err:
            self.logger.exception('{0} - {1}'.format(calling_func, err.args[0]))
            raise

    def UnzipFile(self, sFile, sTarget, bDelDirs):
        if legacy.os.name == 'nt':
            return super().UnzipFile(sFile, sTarget, bDelDirs)
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        self.logger.debug("{0} - sFile: '{1}'".format(calling_func, sFile))
        self.logger.debug("{0} - sTarget: '{1}'".format(calling_func, sTarget))
        self.logger.debug("{0} - bDelDirs: '{1}'".format(calling_func, bDelDirs))
        from scripthost_portable.file_operations import unzip_file
        return unzip_file(sFile, sTarget, bDelDirs)

    def SPFCopy(self, MySrc, MyDest, continueOnError=False, emitConsoleMessages=True, usePyCopy=False):
        if legacy.os.name == 'nt':
            return super().SPFCopy(MySrc, MyDest, continueOnError, emitConsoleMessages, usePyCopy)
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        self.logger.debug("{0} - MySrc: '{1}'".format(calling_func, MySrc))
        self.logger.debug("{0} - MyDest: '{1}'".format(calling_func, MyDest))
        self.logger.debug("{0} - continueOnError: '{1}'".format(calling_func, continueOnError))
        self.logger.debug("{0} - usePyCopy: '{1}'".format(calling_func, usePyCopy))
        MyReturn = 1
        try:
            if emitConsoleMessages is True:
                self.Console('Starting Copy Applet, v3.1 ... {0}'.format(self.DatetimeNow))
            if self.IsEmptyOrNone(MySrc) is True:
                self.Console('  No source path specified. Exiting ...')
                return MyReturn
            if self.IsEmptyOrNone(MyDest) is True:
                self.Console('  No destination path specified. Exiting ...')
                return MyReturn
            if emitConsoleMessages is True:
                self.Console('  Copying {0} to {1}'.format(MySrc, MyDest))
            from scripthost_portable.file_operations import copy_files
            copy_files(MySrc, MyDest)
            MyReturn = 0
            if emitConsoleMessages is True:
                self.ConsoleDoneWithoutTimeStamp()
            return MyReturn
        except Exception as err:
            self.logger.exception('{0} - {1}'.format(calling_func, err.args[0]))
            raise

    def LoadExcel2(self, MyMode, DoContinue, MyCSVFile, ExcelResultFile, MyXLSFile='', MyWorkSheet='', MyVBProc=''):
        if legacy.os.name == 'nt':
            return super().LoadExcel2(MyMode, DoContinue, MyCSVFile, ExcelResultFile, MyXLSFile, MyWorkSheet, MyVBProc)
        calling_func = self.getCallingFuncName(2, self.__class__.__name__)
        self.logger.debug("{0} - MyMode: '{1}'".format(calling_func, MyMode))
        self.logger.debug("{0} - DoContinue: '{1}'".format(calling_func, DoContinue))
        self.logger.debug("{0} - MyCSVFile: '{1}'".format(calling_func, MyCSVFile))
        self.logger.debug("{0} - ExcelResultFile: '{1}'".format(calling_func, ExcelResultFile))
        self.logger.debug("{0} - MyXLSFile: '{1}'".format(calling_func, MyXLSFile))
        self.logger.debug("{0} - MyWorkSheet: '{1}'".format(calling_func, MyWorkSheet))
        self.logger.debug("{0} - MyVBProc: '{1}'".format(calling_func, MyVBProc))
        try:
            from openpyxl import Workbook, load_workbook
            if MyVBProc:
                raise RuntimeError('Excel VBA execution requires the original Windows Excel integration')
            output = legacy.Path(ExcelResultFile or 'SQLPathFinder.xlsx')
            if output.suffix.lower() != '.xlsx' or (MyXLSFile and legacy.Path(MyXLSFile).suffix.lower() != '.xlsx'):
                raise RuntimeError('Portable Excel LOAD/IMPORT supports .xlsx workbooks only')
            files = next(legacy.csv.reader(legacy.StringIO(MyCSVFile), skipinitialspace=True))
            sheets = next(legacy.csv.reader(legacy.StringIO(MyWorkSheet), skipinitialspace=True)) if MyWorkSheet else []
            if MyMode.upper() == 'LOAD':
                files, sheets = ([MyCSVFile], ['Sheet1'])
            elif len(files) != len(sheets):
                raise ValueError('Excel IMPORT requires one worksheet name per CSV file')
            frames = [legacy.pd.read_csv(name, sep=self.GetFileDLM(name), dtype=str, keep_default_na=False, encoding=self.detectFileEncoding(name, readall=True)) for name in files]
            workbook = load_workbook(MyXLSFile) if MyXLSFile else Workbook()
            if not MyXLSFile:
                workbook.remove(workbook.active)
            for frame, name in zip(frames, sheets):
                sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
                sheet.delete_rows(1, sheet.max_row)
                sheet.append(list(frame.columns))
                for row in frame.itertuples(index=False, name=None):
                    sheet.append(list(row))
            workbook.save(output)
            workbook.close()
            return
        except Exception as err:
            self.logger.exception('{0} - {1}'.format(calling_func, err.args[0]))
            if DoContinue is False:
                self.gMyAbort = True
                raise
