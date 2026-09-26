from __future__ import annotations
import shutil
from pathlib import Path
from typing import TYPE_CHECKING
import pandas as pd
from vg2c_new.paths import resolve_path, working_directory_for
from vg2c_new.utilities.base import Utility
from vg2c_new.utilities.csv import CsvUtility
if TYPE_CHECKING:
    from vg2c_new.model import Command
    from vg2c_new.runtime import RuntimeState

class XlsToCsvUtility(Utility):
    def apply(self,command:Command,state:RuntimeState)->None:
        """Rewrite of current XLSToCSV semantics with pandas Excel engines.
        Source: SPSQL3_py/SPFLib/SPFSQL3.py :: XLSToCSVTask. Reference source/commit: vendored ScriptHost at 8ddd5e6463b43834d769057be48041ec657f0f9d.
        Port mode: REWRITE. Preserved: worksheet selection, top/footer skipping and delimited output.
        Amendments: portable pandas/openpyxl/xlrd. Intentionally discarded: COM/helper executable and V1.
        """
        args=[state.substitute(v) for v in command.arguments]
        if len(args)<2: raise ValueError("XLSToCSV requires Excel input and CSV output.")
        source,output=_path(command,state,args[0]),_path(command,state,args[1]); sheet:int|str=0
        if len(args)>2 and args[2].strip(): sheet=int(args[2])-1 if args[2].strip().isdigit() else args[2]
        skiprows=int(args[3]) if len(args)>3 and args[3].isdigit() else 0; skipfooter=int(args[4]) if len(args)>4 and args[4].isdigit() else 0
        frame=pd.read_excel(source,sheet_name=sheet,skiprows=skiprows,skipfooter=skipfooter,dtype=str).fillna("")
        frame.columns=[str(c).replace("\n"," ") for c in frame.columns]; CsvUtility.write_dataframe(frame,output)

class LoadExcelUtility(Utility):
    def apply(self,command:Command,state:RuntimeState)->None:
        """Rewrite of ScriptHost LoadExcel current Version-2 behavior using openpyxl.
        Source: SPSQL3_py/SPFLib/SPFSQL3.py :: LoadExcelTask -> Utilities.LoadExcel2.
        Reference source/commit: vendored ScriptHost at 8ddd5e6463b43834d769057be48041ec657f0f9d.
        Port mode: REWRITE. Preserved: CSV-to-workbook and continue; VERSION 1 syntax flattens to V2 as ScriptHost already does.
        Amendments: pandas/openpyxl replace spfExcelUtility.exe. Intentionally discarded: COM transport/history.
        """
        args=[state.substitute(v) for v in command.arguments]
        if len(args)<2: raise ValueError("LoadExcel requires CSV input and Excel output.")
        version=args[2].strip().upper() if len(args)>2 else "VERSION 2"
        if version not in {"","VERSION 1","VERSION 2"}: raise ValueError(f"Unsupported LoadExcel version selector {version!r}.")
        cont=_yn(args[3]) if len(args)>3 else False
        try:
            frame=CsvUtility.read_dataframe(_path(command,state,args[0])); output=_path(command,state,args[1]); output.parent.mkdir(parents=True,exist_ok=True)
            frame.to_excel(output,index=False,engine="openpyxl")
        except Exception:
            if not cont: raise

class ImportExcelUtility(Utility):
    def apply(self,command:Command,state:RuntimeState)->None:
        """Rewrite of ScriptHost ImportExcel current Version-2 behavior using openpyxl.
        Source: SPSQL3_py/SPFLib/SPFSQL3.py :: ImportExcelTask -> Utilities.LoadExcel2.
        Reference source/commit: vendored ScriptHost at 8ddd5e6463b43834d769057be48041ec657f0f9d.
        Port mode: REWRITE. Preserved: template/result, CSV, worksheet, continue and V1->V2 flattening.
        Amendments: openpyxl replaces spfExcelUtility.exe. Intentionally discarded: VBA and COM.
        """
        args=[state.substitute(v) for v in command.arguments]
        if len(args)<4: raise ValueError("ImportExcel requires template, result workbook, CSV and worksheet.")
        source_book,result_book,csv_name,worksheet=args[:4]; vb_proc=args[5].strip() if len(args)>5 else ""
        version=args[6].strip().upper() if len(args)>6 else "VERSION 2"; cont=_yn(args[7]) if len(args)>7 else False
        if version not in {"","VERSION 1","VERSION 2"}: raise ValueError(f"Unsupported ImportExcel version selector {version!r}.")
        if vb_proc: raise RuntimeError("VBA procedure execution is intentionally not supported in the portable Excel runtime.")
        try:
            from openpyxl import Workbook,load_workbook
            source=_path(command,state,source_book); result=_path(command,state,result_book); result.parent.mkdir(parents=True,exist_ok=True)
            if source.exists(): shutil.copy2(source,result)
            else: Workbook().save(result)
            wb=load_workbook(result)
            if worksheet in wb.sheetnames:
                ws=wb[worksheet]; ws.delete_rows(1,ws.max_row)
            else: ws=wb.create_sheet(worksheet)
            frame=CsvUtility.read_dataframe(_path(command,state,csv_name)); ws.append(list(frame.columns))
            for row in frame.itertuples(index=False,name=None): ws.append(list(row))
            wb.save(result)
        except Exception:
            if not cont: raise

def _path(command:Command,state:RuntimeState,value:str)->Path: return resolve_path(value,state,base=working_directory_for(command.option("WORKDIR"),state))
def _yn(value:str)->bool: return str(value).strip().upper() in {"Y","YES","TRUE","1"}
